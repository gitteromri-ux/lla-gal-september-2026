import openpyxl, datetime as dt, html
SRC='/home/user/workspace/gal/Longevity GAL September 2026.xlsx'
OUT='/home/user/workspace/lla-gal-september-2026/index.html'
wb=openpyxl.load_workbook(SRC,data_only=True)

def fmt(v, nf):
    if v is None or v=='': return ''
    if isinstance(v,dt.datetime): return f'{v.day}.{v.month}.{v.year}'
    if isinstance(v,str): return html.escape(v)
    if isinstance(v,(int,float)):
        if v==0 and nf.count(';')>=2: return '0' if nf.rstrip().endswith(';0') else '–'
        if '%' in nf: return f'{v*100:.1f}%' if '0.0%' in nf else f'{v*100:.0f}%'
        if nf.startswith('0.00'): return f'{v:.2f}'
        if '0.0' in nf: return f'{v:,.1f}'
        return f'{v:,.0f}'
    return html.escape(str(v))

def color(c):
    try:
        rgb=c.fill.start_color.rgb
        if isinstance(rgb,str) and len(rgb)==8 and rgb!='00000000': return '#'+rgb[2:]
    except: pass
    return None
def fcolor(c):
    try:
        rgb=c.font.color.rgb
        if isinstance(rgb,str) and len(rgb)==8: return '#'+rgb[2:]
    except: pass
    return None

def sheet_html(ws, hidden_cols=(), header_rows=(), skip_rows=()):
    merged={}
    for m in ws.merged_cells.ranges:
        merged[(m.min_row,m.min_col)]=(m.max_row-m.min_row+1,m.max_col-m.min_col+1)
        for r in range(m.min_row,m.max_row+1):
            for c in range(m.min_col,m.max_col+1):
                if (r,c)!=(m.min_row,m.min_col): merged[(r,c)]=None
    maxc=ws.max_column
    hidden=set(hidden_cols)|{c for c in range(1,maxc+1) if ws.column_dimensions[openpyxl.utils.get_column_letter(c)].hidden}
    out=['<table dir="ltr">']
    head_rows = header_rows
    for r in range(1,ws.max_row+1):
        if r in skip_rows or ws.row_dimensions[r].hidden: continue
        if head_rows and r==head_rows[0]: out.append('<thead>')
        cells=[]
        empty=True
        for c in range(1,maxc+1):
            if c in hidden: continue
            if (r,c) in merged and merged[(r,c)] is None: continue
            cell=ws.cell(r,c)
            v=fmt(cell.value, cell.number_format or '')
            if v!='': empty=False
            st=[]
            bg=color(cell); fg=fcolor(cell)
            if bg: st.append(f'background:{bg}')
            if fg: st.append(f'color:{fg}')
            if cell.font.bold: st.append('font-weight:700')
            if cell.font.italic: st.append('font-style:italic')
            if cell.font.sz: st.append(f'font-size:{int(cell.font.sz)+8}px')
            al=cell.alignment.horizontal
            if isinstance(cell.value,(int,float)) and not isinstance(cell.value,dt.datetime) and al is None: al='right'
            if al: st.append(f'text-align:{al}')
            span=''
            if (r,c) in merged and merged[(r,c)]:
                rs,cs=merged[(r,c)]
                cs-=len([x for x in range(c,c+cs) if x in hidden])
                span=f' colspan="{cs}" rowspan="{rs}"'
            cells.append(f'<td{span} style="{";".join(st)}">{v}</td>')
        if not empty: out.append('<tr>'+''.join(cells)+'</tr>')
        if head_rows and r==head_rows[-1]: out.append('</thead><tbody>')
    out.append('</tbody></table>' if head_rows else '</table>')
    return '\n'.join(out)

sections=[]
for ws in wb:
    if ws.title=='GAL Sep-26':
        body=sheet_html(ws, header_rows=(4,5), skip_rows=(1,2,3))
    elif ws.title=='Spend by Campaign':
        body=sheet_html(ws, header_rows=(4,), skip_rows=(1,2,3))
    else:
        body=sheet_html(ws)
    sections.append(f'<section><h2>{html.escape(ws.title)}</h2><div class="wrap">{body}</div></section>')

page=f'''<!doctype html>
<html lang="en" dir="ltr">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Longevity Life Academy — GAL Board Report · September 2026</title>
<style>
html,body{{direction:ltr !important;unicode-bidi:embed}}
body{{margin:0;padding:28px 32px 60px;font-family:Calibri,"Segoe UI",Arial,sans-serif;font-size:26px;color:#1F2937;background:#F7F8FB}}
h1{{font-size:44px;color:#0B1F3A;margin:0 0 10px}}
h2{{font-size:34px;color:#0B1F3A;margin:40px 0 14px}}
.bar{{display:flex;gap:14px;flex-wrap:wrap;margin:14px 0 6px}}
.bar a{{display:inline-block;padding:14px 24px;border-radius:12px;background:#0B1F3A;color:#fff;text-decoration:none;font-weight:700;font-size:24px}}
.bar a.alt{{background:#0E7C86}}
.note{{font-size:22px;color:#4B5563;margin:6px 0 18px;line-height:1.4}}
.wrap{{overflow:visible;background:#fff;border:1px solid #D6DBE3;border-radius:12px;box-shadow:0 2px 10px rgba(11,31,58,.06)}}
table{{border-collapse:separate;border-spacing:0;direction:ltr;white-space:nowrap;min-width:100%}}
td{{padding:12px 16px;border-bottom:1px solid #E8EDF5;font-size:26px;text-align:left;vertical-align:middle;line-height:1.25}}
thead td{{position:sticky;top:0;z-index:5;box-shadow:0 2px 0 #0B1F3A}}
thead tr:nth-child(2) td{{top:52px}}
thead tr:first-child td{{height:28px}}
tbody td:first-child, thead td:first-child{{position:sticky;left:0;z-index:6;background:#fff;box-shadow:2px 0 0 #D6DBE3;white-space:normal;min-width:260px;max-width:380px}}
tbody td:first-child[colspan]{{max-width:none;white-space:nowrap}}
thead td:first-child{{z-index:8}}
thead tr:first-child td:first-child{{left:auto;position:sticky}}
tbody tr td:first-child[style*="background"]{{}}
td[colspan]{{text-align:left}}
</style>
</head>
<body>
<h1>Longevity Life Academy — GAL Board Report · September 2026</h1>
<div class="note">Pink rows = weeks 1–2 (1–11.9) exactly as already reported, not touched. Weeks 3–5 (14–28.9) completed from BI By-Day pivot 38 + Meta Ads spend split lead-gen vs e-commerce Blueprint. Layout is fixed left-to-right; dates are in the first column on the left.</div>
<div class="bar">
<a href="Longevity-GAL-September-2026.xlsx">Download Excel</a>
<a class="alt" href="Longevity-GAL-September-2026.pdf">Open PDF</a>
<a class="alt" href="https://github.com/gitteromri-ux/lla-gal-september-2026">GitHub repo</a>
</div>
{''.join(sections)}
<script>
function fixSticky(){{document.querySelectorAll('table').forEach(t=>{{const r=t.querySelectorAll('thead tr');if(r.length>1){{const h=r[0].getBoundingClientRect().height;r[1].querySelectorAll('td').forEach(td=>td.style.top=h+'px');}}}});}}
window.addEventListener('load',fixSticky);window.addEventListener('resize',fixSticky);
</script>
</body></html>'''
open(OUT,'w').write(page)
print('ok', len(page))
